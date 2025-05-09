from openpyxl import load_workbook
from datetime import date, datetime
import pandas as pd
import os
from copy import copy

# copy format cột trước đó
def copy_cell_format(src_cell, tgt_cell):
    if src_cell.has_style:
        if src_cell.font: tgt_cell.font = copy(src_cell.font)
        if src_cell.border: tgt_cell.border = copy(src_cell.border)
        if src_cell.fill: tgt_cell.fill = copy(src_cell.fill)
        if src_cell.number_format: tgt_cell.number_format = src_cell.number_format
        if src_cell.protection: tgt_cell.protection = copy(src_cell.protection)
        if src_cell.alignment: tgt_cell.alignment = copy(src_cell.alignment)

# hiện đang lấy file, insert 1 cột và lưu vào file mới cho dễ debug
def update_report_file():
    file_path = "INPUT DATA/INPUT DATA.xlsx"
    # file_path = "INPUT DATA/INPUT DATA/output_report_2025-05-08.xlsx"
    today = pd.Timestamp(date.today())
    # today = pd.Timestamp(date.today()) - pd.Timedelta(days=1)
    today_str = today.strftime("%Y-%m-%d")

    # Tạo tên file mới
    dir_name, base_name = os.path.split(file_path)
    name_part, ext = os.path.splitext(base_name)
    new_file = os.path.join(dir_name, f"{name_part}_{today_str}{ext}")

    # Đọc dữ liệu kiểm tra ngày
    df = pd.read_excel(file_path, header=None)
    weekdays_row = df.iloc[0]
    dates_row = pd.to_datetime(df.iloc[1], errors='coerce')

    valid_cols = [i for i in df.columns if pd.notna(weekdays_row[i]) and pd.notna(dates_row[i])]
    if not valid_cols:
        print("⚠️ Không tìm thấy cột hợp lệ.")
        return

    last_col = valid_cols[-1]
    last_date = dates_row[last_col].normalize()

    if last_date == today:
        print("✅ Ngày hôm nay đã có trong bảng.")
        return

    insert_at = last_col + 1
    weekday = today.strftime("%a").upper()

    # Load workbook
    wb = load_workbook(file_path)
    ws = wb.active

    ws.insert_cols(insert_at + 1)

    # Copy format từ cột trước đó
    for row in range(1, ws.max_row + 1):
        src_cell = ws.cell(row=row, column=insert_at)
        tgt_cell = ws.cell(row=row, column=insert_at + 1)
        copy_cell_format(src_cell, tgt_cell)

    # Ghi dữ liệu ngày và thứ
    ws.cell(row=1, column=insert_at + 1).value = weekday
    ws.cell(row=2, column=insert_at + 1).value = today

    wb.save(new_file)
    print(f"✅ Đã lưu file mới: {new_file}")


def calculate_summary_stock(file_path_input):
    xls = pd.ExcelFile(file_path_input)
    last_sheet_name = xls.sheet_names[-1]  # last sheet

    # Đọc dữ liệu từ sheet cuối cùng, tên cột bắt đầu từ hàng t2
    df = pd.read_excel(xls, sheet_name=last_sheet_name, header=1) # sai từ đoạn đọc file excel vậy trời

    # Chuẩn hóa tên cột về chữ thường
    df.columns = [str(col).lower() for col in df.columns]

    # Làm sạch dữ liệu: loại bỏ dấu phẩy và chuyển tất cả các giá trị thành số nguyên
    df = df.replace({',': ''}, regex=True)  # Loại bỏ dấu phẩy trong toàn bộ bảng
    df = df.apply(pd.to_numeric, errors='coerce')  # Chuyển các giá trị thành số, lỗi sẽ trở thành NaN
    df = df.dropna()  # Bỏ dòng có giá trị NaN

    # Tìm cột 'thành tiền' (ignore case)
    thanh_tien_col = next((col for col in df.columns if "thành tiền" in col.lower()), None)

    if thanh_tien_col:
        # Tính tổng
        total_thanh_tien = df[thanh_tien_col].sum()
        # parse ngày tháng năm theo tên sheet
        current_year = date.today().year
        full_date_str = f"{last_sheet_name}.{current_year}"
        parsed_date = datetime.strptime(full_date_str, "%d.%m.%Y").date()
        formatted_date = parsed_date.strftime("%d/%m/%Y")

        print(total_thanh_tien, formatted_date)

# update_report_file()
# calculate_summary_stock("INPUT DATA/INPUT DATA/server.HN STOCK/TỒN KHO NGỌC THỤY THÁNG 04.2025.xlsx")
calculate_summary_stock("INPUT DATA/INPUT DATA/server.SG STOCK/2025/THÁNG 4.xlsx")
